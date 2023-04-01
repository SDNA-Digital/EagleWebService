from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import psycopg2


def InportExcel(request, id):
    html = "<html><body>Done! %s </body></html>"

    conn = psycopg2.connect(
        host="localhost",
        database="DPP",
        user="postgres",
        password="3621"
    )

    cursor = conn.cursor()
    query = f"""select ConfigArquivosLocal, ConfigArquivosTabela from ConfigArquivos where ConfigArquivosId = {id}"""

    cursor.execute(query)
    resultados = cursor.fetchall()
    for resultado in resultados:
        Arquivo: object = resultado[0]
        Tabela = resultado[1]

    # print(Arquivo)
    # print(Tabela)

    query = f"""select ConfigArquivosAssociacaoAtt, ConfigArquivosAssociacaoColuna from ConfigArquivosAssociacao where ConfigArquivosId = {id}"""
    cursor.execute(query)
    resultados = cursor.fetchall()
    contador = 0
    Att = []
    Coluna = []
    for resultado in resultados:
        Att.append(resultado[0])
        Coluna.append(resultado[1])
        contador += 1

    listaatt = ''
    contadoraux = 0
    for attini in Att:
        if contadoraux == contador - 1:
            listaatt += attini
            break
        else:
            contadoraux += 1
            listaatt += attini + ','

    # pega os dados do arquivo excel e monta a query de gravação
    # abre o arquivo excel
    wb = load_workbook(Arquivo)
    ws = wb.active
    contalinhas = 0
    for row in ws.rows:
        contalinhas += 1

    valores = []
    contalinhas += 1
    contavalor = 0
    x = 2
    for row in range(x, contalinhas):
        print(valores)
    print(row)
    for i in range(contador):
        j = i + 1
        letter = get_column_letter(j)
        valores.append(ws[letter + str(row)].value)
        formato = ws[letter + str(row)].value
    print(valores)
    contadoraux = 0
    listavalor = ''
    for valorcampo in valores:
        tipo = str(type(valorcampo))
        tipocompara = "<class 'datetime.datetime'>"
        concatena = valorcampo

        if tipo == tipocompara:
            concatena = ''
            concatena = str(valorcampo)
            concatena = "to_date('" + concatena[8:10] + "/" + concatena[5:7] + "/" + concatena[0:4] + "', 'DD/MM/YYYY')"

        if contadoraux == contador - 1:
            if tipo == tipocompara:
                listavalor += concatena
            else:
                if concatena is not None:
                    listavalor += "'" + concatena + "'"
                break
        else:
            if tipo == tipocompara:
                listavalor += concatena + ','
            else:
                listavalor += "'" + concatena + "'" + ','
            contadoraux += 1

    montaquery = 'INSERT INTO ' + Tabela + ' (' + listaatt + ') Values(' + listavalor + ')'
    print(montaquery)
    query = montaquery
    cursor.execute(query)
    conn.commit()
    x += 1
    valores = []


    return HttpResponse(html)